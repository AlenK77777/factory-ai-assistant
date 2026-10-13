"""Dashboard numbers, the offline HTML page and the Excel workbook. No live model."""

import pytest
import zipfile
from datetime import date, datetime, timezone
from io import BytesIO
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart

from chat.event_log import EventLog
from dashboard.data import ai_report, production_report
from dashboard.render import render_excel, render_html


def sample_production():
    today = date(2026, 10, 13)
    orders = [
        {"order_id": "PO-1", "customer_sector": "Medical", "status": "On Hold", "due_date": date(2026, 10, 1), "quantity": 100},
        {"order_id": "PO-2", "customer_sector": "Medical", "status": "Shipped", "due_date": date(2026, 9, 1), "quantity": 50},
        {"order_id": "PO-3", "customer_sector": "Aerospace", "status": "In Production", "due_date": date(2026, 12, 1), "quantity": 10},
    ]
    defects = [
        {"detected_date": date(2026, 10, 1), "station": "AOI", "defect_type": "Open", "quantity": 4, "customer_sector": "Medical"},
        {"detected_date": date(2026, 10, 2), "station": "AOI", "defect_type": "Short", "quantity": 1, "customer_sector": "Medical"},
        {"detected_date": date(2026, 9, 1), "station": "Plating", "defect_type": "Open", "quantity": 3, "customer_sector": "Aerospace"},
    ]
    runs = [
        {
            "file_name": "defects.xlsx",
            "started_at": datetime(2026, 10, 12, 8, 0, tzinfo=timezone.utc),
            "rows_loaded": 8,
            "rows_rejected": 2,
        }
    ]
    return production_report(orders, defects, runs, today)


def test_production_kpis_late_orders_and_weeks():
    report = sample_production()
    assert report["kpis"]["total_orders"] == 3
    assert report["kpis"]["on_hold"] == 1
    # PO-2 is past due but Shipped, so it is not late. PO-1 is.
    assert report["kpis"]["late"] == 1
    assert report["late_orders"][0]["order_id"] == "PO-1"
    assert report["late_orders"][0]["days_late"] == 12
    assert report["kpis"]["defective_units"] == 8
    assert report["kpis"]["ordered_units"] == 160
    assert report["kpis"]["defect_rate"] == round(800 / 160, 1)
    statuses = {row["name"]: row["count"] for row in report["orders_by_status"]}
    assert statuses["On Hold"] == 1
    assert statuses["Shipped"] == 1
    stations = {row["name"]: row["units"] for row in report["defects_by_station"]}
    assert stations["AOI"] == 5
    assert stations["Plating"] == 3
    assert report["defects_by_sector"][0]["name"] == "Medical"
    assert len(report["weekly_defects"]) == 2
    assert report["weekly_defects"][0]["week"] < report["weekly_defects"][1]["week"]
    assert report["weekly_defects"][0]["partial"] is True
    assert report["weekly_defects"][0]["label"].endswith(" partial")
    assert report["weekly_defects"][1]["partial"] is False
    assert report["ingest_runs"][0]["rows_rejected"] == 2
    assert report["ingest_runs"][0]["file_name"] == "defects.xlsx"


def test_ai_report_from_question_events(tmp_path):
    log = EventLog(tmp_path / "events.log")
    log.write(
        "INFO",
        "chat",
        "question",
        language="en",
        model="qwen3:14b",
        outcome="not_found",
        source="user",
        latency_s=1.5,
        question="What is the weather?",
        tools=[{"name": "search_instructions", "arguments": {"query": "What is the weather?"}}],
    )
    log.write(
        "INFO",
        "chat",
        "question",
        language="ru",
        model="qwen3:14b",
        outcome="found",
        source="user",
        latency_s=3.5,
        question="status",
        tools=[{"name": "get_order_status", "arguments": {"order_id": "PO-1"}}],
    )
    log.write("INFO", "server", "health", ollama="ok")
    report = ai_report(log.read_events())
    assert report["outcomes"] == {"found": 1, "not_found": 1, "error": 0}
    assert report["by_language"] == [{"name": "en", "count": 1}, {"name": "ru", "count": 1}]
    names = {row["name"]: row["count"] for row in report["tools"]}
    assert names == {"search_instructions": 1, "get_order_status": 1}
    assert report["latency_by_model"] == [{"model": "qwen3:14b", "count": 2, "median_s": 2.5, "slowest_s": 3.5}]
    assert report["unanswered"] == [
        {
            "timestamp": report["unanswered"][0]["timestamp"],
            "language": "en",
            "model": "qwen3:14b",
            "question": "What is the weather?",
        }
    ]
    assert len(report["questions_per_day"]) == 1


def test_legacy_and_smoke_are_excluded_unless_test_traffic_is_included(tmp_path):
    log = EventLog(tmp_path / "events.log")
    log.write("INFO", "chat", "question", language="en", outcome="found", source="user", question="from the page")
    log.write("INFO", "chat", "question", language="en", outcome="found", source="smoke", question="from the smoke test")
    log.write("INFO", "chat", "question", language="en", outcome="not_found", question="written before source existed")
    events = log.read_events()
    quiet = ai_report(events)
    assert quiet["outcomes"] == {"found": 1, "not_found": 0, "error": 0}
    assert quiet["unanswered"] == []
    full = ai_report(events, include_test=True)
    assert full["outcomes"] == {"found": 2, "not_found": 1, "error": 0}
    assert len(full["unanswered"]) == 1


def test_html_is_self_contained_svg():
    report = {
        "generated_at": "2026-10-13T15:00:00+03:00",
        "production": sample_production(),
        "ai": ai_report([]),
    }
    document = render_html(report)
    lowered = document.lower()
    assert "Generated at 2026-10-13T15:00:00+03:00" in document
    assert "Test traffic included: no" in document
    assert "found an answer (correctness not checked)" in document
    assert "Answered" not in document
    assert "<svg" in lowered
    assert "<script" not in lowered
    assert "cdn" not in lowered
    assert "http://" not in lowered
    assert "https://" not in lowered
    assert "Total orders" in document
    assert ">3<" in document
    assert "PO-1" in document
    assert " partial" in document
    assert "Station" in document
    assert "Defective units" in document
    assert "No data." in document


def test_excel_has_a_sheet_per_section_and_native_charts():
    report = {
        "generated_at": "2026-10-13T15:00:00+03:00",
        "production": sample_production(),
        "ai": {
            "questions_per_day": [{"day": "2026-10-13", "count": 2}],
            "by_language": [{"name": "en", "count": 2}],
            "outcomes": {"found": 1, "not_found": 1, "error": 0},
            "tools": [{"name": "get_order_status", "count": 1}],
            "latency_by_model": [{"model": "qwen3:14b", "count": 2, "median_s": 2.5, "slowest_s": 3.5}],
            "unanswered": [{"timestamp": "2026-10-13T15:00:00+03:00", "language": "en", "model": "qwen3:14b", "question": "weather"}],
        },
    }
    workbook = load_workbook(BytesIO(render_excel(report)))
    assert workbook.sheetnames[0] == "Overview"
    for name in (
        "Orders by status",
        "Defects by station",
        "Defects by type",
        "Defects by sector",
        "Weekly defects",
        "Late orders",
        "Ingest runs",
        "Questions per day",
        "Questions by language",
        "Outcomes",
        "Tool usage",
        "Latency by model",
        "Unanswered",
    ):
        assert name in workbook.sheetnames
    assert workbook["Overview"]["B1"].value == datetime(2026, 10, 13, 15, 0)
    assert workbook["Overview"]["B12"].value == "=IF(B11=0,0,B10/B11)"
    assert workbook["Overview"]["B12"].number_format == "0.0%"
    assert workbook["Overview"]["B2"].value == "no"
    assert workbook["Outcomes"]["A2"].value == "found an answer (correctness not checked)"
    orders = workbook["Orders by status"]
    assert orders["A2"].value == "In Production"
    assert orders["B2"].value == 1
    assert orders["A3"].value == "On Hold"
    assert len(orders._charts) == 1
    assert isinstance(orders._charts[0], BarChart)
    assert orders._charts[0].series
    weekly = workbook["Weekly defects"]
    assert isinstance(weekly._charts[0], LineChart)
    assert weekly._charts[0].series
    latency = workbook["Latency by model"]
    assert latency["B2"].value == 2.5
    assert latency["C2"].value == 3.5
    assert len(latency._charts[0].series) == 2
    assert workbook["Late orders"]._charts == []
    assert workbook["Unanswered"]["D2"].value == "weather"
    assert workbook["Ingest runs"]["D2"].value == 2
    assert isinstance(workbook["Late orders"]["D2"].value, date)
    assert isinstance(workbook["Ingest runs"]["B2"].value, datetime)
    assert isinstance(workbook["Questions per day"]["A2"].value, date)


CHART_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}


def test_excel_chart_xml_keeps_axes_and_text_categories():
    report = {
        "generated_at": "2026-10-13T15:00:00+03:00",
        "production": sample_production(),
        "ai": {
            "questions_per_day": [{"day": "2026-10-13", "count": 2}],
            "by_language": [{"name": "en", "count": 2}],
            "outcomes": {"found": 1, "not_found": 1, "error": 0},
            "tools": [{"name": "get_order_status", "count": 1}],
            "latency_by_model": [{"model": "qwen3:14b", "count": 2, "median_s": 2.5, "slowest_s": 3.5}],
            "unanswered": [],
        },
    }
    blob = render_excel(report)
    with zipfile.ZipFile(BytesIO(blob)) as book:
        names = [name for name in book.namelist() if name.startswith("xl/charts/chart")]
        assert names
        for name in names:
            root = ET.fromstring(book.read(name))
            axes = root.findall(".//c:catAx", CHART_NS) + root.findall(".//c:valAx", CHART_NS)
            assert axes
            for axis in axes:
                deleted = axis.find("c:delete", CHART_NS)
                assert deleted is None or deleted.get("val") in ("0", "false")
                assert axis.find("c:title", CHART_NS) is not None
            for category in root.findall(".//c:cat", CHART_NS):
                assert category.find("c:strRef", CHART_NS) is not None
                assert category.find("c:numRef", CHART_NS) is None


@pytest.mark.integration
def test_live_database_report_when_available():
    """Uses the read-only role when the database is up. Skipped in CI."""
    from dashboard.data import fetch_production

    report = fetch_production()
    if report.get("error"):
        pytest.skip(report["error"])
    assert report["kpis"]["total_orders"] >= 1
    assert report["kpis"]["late"] <= report["kpis"]["total_orders"]
    assert len(report["late_orders"]) <= 10
    assert report["kpis"]["defect_rate"] >= 0
