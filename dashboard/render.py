"""Self-contained dashboard HTML (inline SVG) and an Excel workbook with real charts."""

from __future__ import annotations

import html
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

from dashboard.data import fetch_production, load_ai_report

LANGUAGE_NAMES = {"en": "English", "ru": "Russian", "he": "Hebrew"}


def generated_stamp() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def snapshot(today=None) -> dict:
    """Current production and assistant numbers, plus the time they were read."""
    return {
        "generated_at": generated_stamp(),
        "production": fetch_production(today),
        "ai": load_ai_report(),
    }


def _num(value) -> str:
    if isinstance(value, float) and not value.is_integer():
        return f"{value:.1f}"
    return str(int(value) if isinstance(value, float) else value)


def _esc(value) -> str:
    return html.escape(str(value), quote=True)


def svg_bars(title: str, rows: list[tuple[str, float]]) -> str:
    """Horizontal bar chart. Empty input is a short sentence, not an empty picture."""
    if not rows:
        return f"<h3>{_esc(title)}</h3><p>No data.</p>"
    row_h = 28
    label_w = 168
    chart_w = 380
    width = label_w + chart_w + 64
    height = 12 + len(rows) * row_h
    peak = max(value for _, value in rows) or 1
    parts = [
        f"<h3>{_esc(title)}</h3>",
        f'<svg viewBox="0 0 {width} {height}" width="100%" role="img">',
        f"<title>{_esc(title)}</title>",
    ]
    for index, (label, value) in enumerate(rows):
        y = 6 + index * row_h
        bar = chart_w * (float(value) / peak)
        shown = label if len(label) <= 24 else label[:23] + "..."
        parts.append(f'<text x="0" y="{y + 16}" font-size="13" fill="#1c2430">{_esc(shown)}</text>')
        parts.append(f'<rect x="{label_w}" y="{y}" width="{bar:.1f}" height="18" fill="#1f4e79"></rect>')
        parts.append(
            f'<text x="{label_w + bar + 6:.1f}" y="{y + 14}" font-size="12" fill="#1c2430">{_esc(_num(value))}</text>'
        )
    parts.append("</svg>")
    return "\n".join(parts)


def svg_line(title: str, rows: list[tuple[str, float]]) -> str:
    if not rows:
        return f"<h3>{_esc(title)}</h3><p>No data.</p>"
    width, height = 640, 220
    left, right, top, bottom = 40, 16, 16, 36
    peak = max(value for _, value in rows) or 1
    count = len(rows)

    def point(index: int, value: float) -> tuple[float, float]:
        span = width - left - right
        x = left if count == 1 else left + span * index / (count - 1)
        y = top + (height - top - bottom) * (1 - float(value) / peak)
        return x, y

    coords = [point(index, value) for index, (_, value) in enumerate(rows)]
    polyline = " ".join(f"{x:.1f},{y:.1f}" for x, y in coords)
    parts = [
        f"<h3>{_esc(title)}</h3>",
        f'<svg viewBox="0 0 {width} {height}" width="100%" role="img">',
        f"<title>{_esc(title)}</title>",
        f'<line x1="{left}" y1="{top}" x2="{left}" y2="{height - bottom}" stroke="#d5dbe3"/>',
        f'<line x1="{left}" y1="{height - bottom}" x2="{width - right}" y2="{height - bottom}" stroke="#d5dbe3"/>',
        f'<polyline fill="none" stroke="#1f4e79" stroke-width="2" points="{polyline}"/>',
    ]
    for index, ((label, value), (x, y)) in enumerate(zip(rows, coords)):
        parts.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3" fill="#1f4e79"/>')
        if count <= 12 or index in (0, count - 1) or index == count // 2:
            parts.append(
                f'<text x="{x:.1f}" y="{height - 12}" font-size="11" text-anchor="middle" fill="#3d4b5c">{_esc(label)}</text>'
            )
    parts.append(f'<text x="{left}" y="12" font-size="11" fill="#5c6b7a">max {_esc(_num(peak))}</text>')
    parts.append("</svg>")
    return "\n".join(parts)


def svg_grouped(title: str, categories: list[str], series: list[tuple[str, list[float]]]) -> str:
    """Two bars per category. Used for median and slowest latency."""
    if not categories or not series:
        return f"<h3>{_esc(title)}</h3><p>No data.</p>"
    row_h = 36
    label_w = 180
    chart_w = 360
    width = label_w + chart_w + 16
    height = 28 + len(categories) * row_h
    peak = max((value for _, values in series for value in values), default=1) or 1
    colors = ("#1f4e79", "#a15c07")
    parts = [
        f"<h3>{_esc(title)}</h3>",
        f'<svg viewBox="0 0 {width} {height}" width="100%" role="img">',
        f"<title>{_esc(title)}</title>",
    ]
    legend = "  ".join(name for name, _values in series)
    parts.append(f'<text x="{label_w}" y="14" font-size="12" fill="#5c6b7a">{_esc(legend)}</text>')
    for index, label in enumerate(categories):
        y = 24 + index * row_h
        shown = label if len(label) <= 26 else label[:25] + "..."
        parts.append(f'<text x="0" y="{y + 16}" font-size="13" fill="#1c2430">{_esc(shown)}</text>')
        for series_index, (_name, values) in enumerate(series):
            value = values[index] if index < len(values) else 0
            bar = chart_w * (float(value) / peak)
            bar_y = y + series_index * 14
            color = colors[series_index % len(colors)]
            parts.append(f'<rect x="{label_w}" y="{bar_y}" width="{bar:.1f}" height="12" fill="{color}"></rect>')
    parts.append("</svg>")
    return "\n".join(parts)


def html_table(headers: list[str], rows: list[list]) -> str:
    if not rows:
        return "<p>No data.</p>"
    head = "".join(f"<th>{_esc(header)}</th>" for header in headers)
    body = []
    for row in rows:
        cells = "".join(f"<td>{_esc(cell)}</td>" for cell in row)
        body.append(f"<tr>{cells}</tr>")
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def _tiles(kpis: dict) -> str:
    items = [
        ("Total orders", kpis["total_orders"]),
        ("Late", kpis["late"]),
        ("On hold", kpis["on_hold"]),
        ("Defective units", kpis["defective_units"]),
        ("Defect rate", f"{kpis['defect_rate']:.1f}%"),
    ]
    boxes = []
    for label, value in items:
        boxes.append(f'<div class="tile"><span>{_esc(label)}</span><strong>{_esc(value)}</strong></div>')
    return '<div class="tiles">' + "".join(boxes) + "</div>"


def render_html(report: dict) -> str:
    """A full HTML document. Charts are inline SVG. No external files or CDN."""
    production = report["production"]
    ai = report["ai"]
    kpis = production["kpis"]
    outcomes = ai["outcomes"]
    total_questions = sum(outcomes.values())
    share_rows = []
    for key, label in (("found", "Answered"), ("not_found", "Not found"), ("error", "Errors")):
        count = outcomes[key]
        percent = round(100.0 * count / total_questions, 1) if total_questions else 0.0
        share_rows.append((f"{label} ({percent:.1f}%)", count))
    language_rows = [(LANGUAGE_NAMES.get(row["name"], row["name"]), row["count"]) for row in ai["by_language"]]
    sections = [
        "<h2>Production</h2>",
        f"<p class='note'>Defect rate is defective units divided by boards ordered ({kpis['ordered_units']}). Late means the due date is before today and the order is not Shipped.</p>",
        _tiles(kpis),
        svg_bars("Orders by status", [(row["name"], row["count"]) for row in production["orders_by_status"]]),
        svg_bars("Defects by station (units)", [(row["name"], row["units"]) for row in production["defects_by_station"]]),
        svg_bars("Defects by type (units)", [(row["name"], row["units"]) for row in production["defects_by_type"]]),
        svg_bars("Defects by customer sector (units)", [(row["name"], row["units"]) for row in production["defects_by_sector"]]),
        svg_line("Weekly defect trend (units)", [(row["week"], row["units"]) for row in production["weekly_defects"]]),
        "<h3>Top 10 late orders</h3>",
        html_table(
            ["Order", "Sector", "Status", "Due", "Days late"],
            [[row["order_id"], row["customer_sector"], row["status"], row["due_date"], row["days_late"]] for row in production["late_orders"]],
        ),
        "<h3>Last Excel ingestion runs</h3>",
        html_table(
            ["File", "Started", "Loaded", "Rejected"],
            [[row["file_name"], row["started_at"], row["rows_loaded"], row["rows_rejected"]] for row in production["ingest_runs"]],
        ),
        "<h2>AI assistant</h2>",
        "<p class='note'>Counted from the event log. A question is Not found when the assistant had nothing relevant to answer with.</p>",
        svg_line("Questions per day", [(row["day"], row["count"]) for row in ai["questions_per_day"]]),
        svg_bars("Questions by language", language_rows),
        svg_bars("Answered, not found, errors", share_rows),
        svg_bars("Tool usage", [(row["name"], row["count"]) for row in ai["tools"]]),
        svg_grouped(
            "Latency by model (seconds)",
            [row["model"] for row in ai["latency_by_model"]],
            [
                ("Median", [row["median_s"] for row in ai["latency_by_model"]]),
                ("Slowest", [row["slowest_s"] for row in ai["latency_by_model"]]),
            ],
        ),
        html_table(
            ["Model", "Questions", "Median seconds", "Slowest seconds"],
            [[row["model"], row["count"], row["median_s"], row["slowest_s"]] for row in ai["latency_by_model"]],
        ),
        "<h3>Questions with no answer found</h3>",
        html_table(
            ["Time", "Language", "Model", "Question"],
            [[row["timestamp"], row["language"], row["model"], row["question"]] for row in ai["unanswered"]],
        ),
    ]
    error = ""
    if production.get("error"):
        error = f"<p class='error'>{_esc(production['error'])}</p>"
    style = """
    body { margin: 0; font: 15px/1.4 "Segoe UI", sans-serif; color: #1c2430; background: #fff; }
    main { padding: 16px 18px 32px; }
    h1 { font-size: 22px; margin: 0 0 4px; }
    h2 { font-size: 18px; margin: 22px 0 8px; border-top: 1px solid #d5dbe3; padding-top: 14px; }
    h3 { font-size: 15px; margin: 16px 0 6px; }
    .when, .note { color: #5c6b7a; margin: 0 0 8px; }
    .error { background: #fde8e6; border: 1px solid #e3b2ac; color: #6d2018; padding: 8px 10px; }
    .tiles { display: flex; flex-wrap: wrap; gap: 8px; margin: 8px 0 12px; }
    .tile { background: #f4f7fb; border: 1px solid #d5dbe3; border-radius: 8px; padding: 8px 10px; min-width: 110px; }
    .tile span { display: block; color: #5c6b7a; font-size: 12px; }
    .tile strong { font-size: 22px; }
    table { border-collapse: collapse; width: 100%; margin: 0 0 8px; }
    th, td { border: 1px solid #d5dbe3; padding: 4px 6px; text-align: left; vertical-align: top; }
    th { background: #f4f7fb; }
    svg { display: block; margin: 0 0 8px; }
    """
    return (
        "<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>"
        "<title>Factory dashboard</title><style>"
        + style
        + "</style></head><body><main>"
        + "<h1>Factory dashboard</h1>"
        + f"<p class='when'>Generated at {_esc(report['generated_at'])}</p>"
        + "<p class='note'>Synthetic data. This file has no external images, scripts or fonts.</p>"
        + error
        + "\n".join(sections)
        + "</main></body></html>"
    )


def _add_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    chart: str | None = None,
    chart_columns: int | None = None,
) -> None:
    """One sheet: a table, and a native Excel chart bound to those cells."""
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(12, min(42, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
    if not chart or not rows:
        return
    drawing = BarChart() if chart == "bar" else LineChart()
    if chart == "bar":
        drawing.type = "col"
    drawing.title = title
    drawing.y_axis.title = None
    drawing.x_axis.title = None
    drawing.style = 10
    last_col = len(headers) if chart_columns is None else 1 + chart_columns
    data = Reference(sheet, min_col=2, max_col=last_col, min_row=1, max_row=1 + len(rows))
    categories = Reference(sheet, min_col=1, min_row=2, max_row=1 + len(rows))
    drawing.add_data(data, titles_from_data=True)
    drawing.set_categories(categories)
    if chart == "bar":
        drawing.shape = 4
    drawing.width = 18
    drawing.height = 8
    if len(headers) == 2:
        drawing.legend = None
    sheet.add_chart(drawing, "A" + str(len(rows) + 3))


def render_excel(report: dict) -> bytes:
    """Workbook with one sheet per section. Charts point at the cells, so the numbers stay editable."""
    production = report["production"]
    ai = report["ai"]
    kpis = production["kpis"]
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview["A1"] = "Generated at"
    overview["B1"] = report["generated_at"]
    overview["A2"] = "Note"
    overview["B2"] = "Synthetic data. Defect rate is defective units / boards ordered. Late means due before today and not Shipped."
    if production.get("error"):
        overview["A3"] = "Production error"
        overview["B3"] = production["error"]
    overview.append([])
    overview.append(["KPI", "Value"])
    for label, value in (
        ("Total orders", kpis["total_orders"]),
        ("Late", kpis["late"]),
        ("On hold", kpis["on_hold"]),
        ("Defective units", kpis["defective_units"]),
        ("Boards ordered", kpis["ordered_units"]),
        ("Defect rate percent", kpis["defect_rate"]),
    ):
        overview.append([label, value])
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 80

    _add_sheet(
        workbook,
        "Orders by status",
        ["Status", "Orders"],
        [[row["name"], row["count"]] for row in production["orders_by_status"]],
        "bar",
    )
    _add_sheet(
        workbook,
        "Defects by station",
        ["Station", "Units"],
        [[row["name"], row["units"]] for row in production["defects_by_station"]],
        "bar",
    )
    _add_sheet(
        workbook,
        "Defects by type",
        ["Defect type", "Units"],
        [[row["name"], row["units"]] for row in production["defects_by_type"]],
        "bar",
    )
    _add_sheet(
        workbook,
        "Defects by sector",
        ["Customer sector", "Units"],
        [[row["name"], row["units"]] for row in production["defects_by_sector"]],
        "bar",
    )
    _add_sheet(
        workbook,
        "Weekly defects",
        ["Week", "Units"],
        [[row["week"], row["units"]] for row in production["weekly_defects"]],
        "line",
    )
    _add_sheet(
        workbook,
        "Late orders",
        ["Order", "Sector", "Status", "Due", "Days late"],
        [[row["order_id"], row["customer_sector"], row["status"], row["due_date"], row["days_late"]] for row in production["late_orders"]],
        None,
    )
    _add_sheet(
        workbook,
        "Ingest runs",
        ["File", "Started", "Loaded", "Rejected"],
        [[row["file_name"], row["started_at"], row["rows_loaded"], row["rows_rejected"]] for row in production["ingest_runs"]],
        None,
    )
    total = sum(ai["outcomes"].values())
    outcome_rows = []
    for key, label in (("found", "Answered"), ("not_found", "Not found"), ("error", "Errors")):
        count = ai["outcomes"][key]
        percent = round(100.0 * count / total, 1) if total else 0.0
        outcome_rows.append([label, count, percent])
    _add_sheet(workbook, "Questions per day", ["Day", "Questions"], [[row["day"], row["count"]] for row in ai["questions_per_day"]], "line")
    _add_sheet(
        workbook,
        "Questions by language",
        ["Language", "Questions"],
        [[LANGUAGE_NAMES.get(row["name"], row["name"]), row["count"]] for row in ai["by_language"]],
        "bar",
    )
    _add_sheet(workbook, "Outcomes", ["Outcome", "Questions", "Percent"], outcome_rows, "bar", chart_columns=1)
    _add_sheet(workbook, "Tool usage", ["Tool", "Calls"], [[row["name"], row["count"]] for row in ai["tools"]], "bar")
    _add_sheet(
        workbook,
        "Latency by model",
        ["Model", "Median seconds", "Slowest seconds"],
        [[row["model"], row["median_s"], row["slowest_s"]] for row in ai["latency_by_model"]],
        "bar",
    )
    _add_sheet(
        workbook,
        "Unanswered",
        ["Time", "Language", "Model", "Question"],
        [[row["timestamp"], row["language"], row["model"], row["question"]] for row in ai["unanswered"]],
        None,
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_html(today=None) -> str:
    return render_html(snapshot(today))


def build_excel(today=None) -> tuple[str, bytes]:
    report = snapshot(today)
    return report["generated_at"], render_excel(report)
